import csv
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
CSV_PATH = ROOT / "PLC_Tags_Palletizer.csv"
XLSX_PATH = ROOT / "PLC_Tags_Palletizer.xlsx"

cols = ["Name", "DataType", "LogicalAddress"]

rows = []
with CSV_PATH.open(newline="", encoding="utf-8") as f:
    reader = csv.DictReader(f, delimiter=";")
    for r in reader:
        rows.append({c: r[c] for c in cols})

wb = Workbook()
ws = wb.active
ws.title = "PLC_Tags"

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="1F4E79")
thin = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

ws.append(cols)
for col in range(1, 4):
    cell = ws.cell(1, col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

for r in rows:
    ws.append([r["Name"], r["DataType"], r["LogicalAddress"]])
    for col in range(1, 4):
        ws.cell(ws.max_row, col).border = thin

ws.column_dimensions["A"].width = 32
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 16
ws.auto_filter.ref = f"A1:C{ws.max_row}"
ws.freeze_panes = "A2"

wb.save(XLSX_PATH)
print(f"Wrote {XLSX_PATH} ({len(rows)} tags)")
