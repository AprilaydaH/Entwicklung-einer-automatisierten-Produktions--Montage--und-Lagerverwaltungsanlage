# -*- coding: utf-8 -*-
"""15-Minuten-Präsentation nach Kursregeln: 15–18 Folien."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "docs"
    / "01_Projektgrundlagen"
    / "Praesentation_Abschlussprojekt_15min.pptx"
)

NAVY = RGBColor(0x1B, 0x36, 0x5D)
TEAL = RGBColor(0x00, 0x7A, 0x99)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x4A, 0x4A, 0x4A)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL = 18


def set_run(run, size=18, bold=False, color=DARK):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, l, t, w, h, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)
    return box


def add_bar(slide, color=NAVY, height=Inches(0.12)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def add_footer(slide, page):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, SLIDE_H - Inches(0.38), SLIDE_W, Inches(0.38)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_textbox(
        slide,
        Inches(0.4),
        SLIDE_H - Inches(0.36),
        Inches(10.2),
        Inches(0.32),
        "Dereje Hailemariam  ·  Abschlussprojekt Weiterbildung  ·  Berlin",
        size=11,
        color=WHITE,
    )
    add_textbox(
        slide,
        Inches(11.4),
        SLIDE_H - Inches(0.36),
        Inches(1.6),
        Inches(0.32),
        f"{page} / {TOTAL}",
        size=11,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
    )


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def title_block(slide, title, subtitle=None):
    add_textbox(slide, Inches(0.5), Inches(0.28), Inches(12.3), Inches(0.5), title, size=26, bold=True, color=NAVY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.84), Inches(2.0), Inches(0.07))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()
    if subtitle:
        add_textbox(slide, Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.32), subtitle, size=14, color=GRAY)


def card(slide, l, t, w, h, heading, body, fill=LIGHT, head_color=NAVY):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.12)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = heading
    set_run(r, size=16, bold=True, color=head_color)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = body
    set_run(r2, size=13, color=DARK)


def bullets(slide, l, t, w, h, items, size=16):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(7)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=DARK)


def table_fill(table, headers, rows, col_widths):
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                set_run(r, size=13, bold=True, color=WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    set_run(run, size=13, bold=(c == 0), color=DARK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if r % 2 else WHITE


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1 Deckblatt
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), SLIDE_H)
    acc.fill.solid()
    acc.fill.fore_color.rgb = TEAL
    acc.line.fill.background()
    add_textbox(s, Inches(0.7), Inches(1.15), Inches(12), Inches(0.35), "Willkommen  ·  Abschlusspräsentation", size=16, color=TEAL)
    add_textbox(
        s,
        Inches(0.7),
        Inches(1.65),
        Inches(12),
        Inches(1.7),
        "Entwicklung einer automatisierten\nProduktions-, Montage- und Lagerverwaltungsanlage",
        size=30,
        bold=True,
        color=WHITE,
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(3.6),
        Inches(12),
        Inches(0.7),
        "RFID-gestützte Produktverfolgung  ·  TIA Portal V20  ·  Factory I/O",
        size=18,
        color=WHITE,
    )
    add_textbox(
        s,
        Inches(0.7),
        Inches(4.7),
        Inches(12),
        Inches(1.5),
        "Name:           Dereje Hailemariam\n"
        "Weiterbildung:  Abschlussprojekt Automatisierungstechnik, Berlin\n"
        "Datum:          19. September 2026\n"
        "Dauer:          15 Minuten",
        size=18,
        color=WHITE,
    )
    notes(s, "~40 s: Begrüßung, Name, Weiterbildung, Datum, Projekttitel. Nicht in Technik gehen.")

    # 2 Gliederung
    s = blank(prs)
    add_bar(s)
    title_block(s, "Inhaltsverzeichnis", "15 Minuten  ·  18 Folien")
    left = [
        "1.  Projektziel",
        "2.  Projektbeschreibung",
        "3.  Soll–Ist-Vergleich",
        "4.  Projektumfang",
        "5.  Kosten",
        "6.  Wirtschaftlichkeit",
        "7.  Zeit",
        "8.  Herausforderungen und Lösungen",
    ]
    right = [
        "9.  Projektdurchführung",
        "      – Anlage und Linien",
        "      – RFID und Hochregal",
        "      – Software und HMI",
        "10. Ausblick / Erweiterungen",
        "11. Fazit — was ich gelernt habe",
        "12. Ende / Fragen",
    ]
    bullets(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(5.2), left, size=18)
    bullets(s, Inches(7.0), Inches(1.4), Inches(5.8), Inches(5.2), right, size=18)
    add_footer(s, 2)
    notes(s, "~30 s: Gliederung vorlesen. Danach strikt dieser Reihenfolge folgen.")

    # 3 Projektziel
    s = blank(prs)
    add_bar(s)
    title_block(s, "Projektziel", "Was am Ende funktionieren soll")
    bullets(
        s,
        Inches(0.5),
        Inches(1.4),
        Inches(12.2),
        Inches(5.4),
        [
            "Eine durchgängige Fertigungs- und Lagerlinie in TIA Portal und Factory I/O",
            "Zwei Materiallinien: Metall (A) und Kunststoff (B), jeweils Base und Deckel",
            "RFID-Verfolgung: Artikel, Material, Typ, Code vom Band bis ins Fach",
            "Montage (2-axis P&P), Palettierung, automatische Ein- und Auslagerung",
            "Eine SPS (CPU 1518F), ein HMI (TP2200), Manual und Automatik",
            "54 Lagerfächer adressierbar, Zustände speicherbar, modular erweiterbar",
            "Dokumentiert: Lastenheft, Pflichtenheft, SCL, Tag-Listen",
        ],
        size=17,
    )
    add_footer(s, 3)
    notes(s, "~1 min: Ziel ist die Verbindung der Teilprozesse, nicht eine einzelne Station.")

    # 4 Beschreibung
    s = blank(prs)
    add_bar(s)
    title_block(s, "Projektbeschreibung", "Simulierte Factory, eine Steuerung")
    card(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(2.35), "Anlage", "Factory I/O: Roh, CNC, Bänder, Vision,\n2-axis / 3-axis Pick & Place, Palettierer,\nzwei Hochregale à 54 Fächer")
    card(s, Inches(6.8), Inches(1.4), Inches(6.1), Inches(2.35), "Steuerung", "TIA Portal V20\nCPU 1518F-4 PN/DP (Failsafe)\nOB1: 29 Netzwerke, SCL + FUP")
    card(s, Inches(0.5), Inches(3.95), Inches(6.1), Inches(2.35), "Identifikation", "Vision erkennt Farbe/Typ\nRFID schreibt den Tag an der Palette\nGate: ohne gültigen Tag kein Auto-Einlagern")
    card(s, Inches(6.8), Inches(3.95), Inches(6.1), Inches(2.35), "Lager", "W1 Kunststoff Zone 5B NW 29\nW2 Metall Zone 5A NW 28\ngetrennte DBs, Merker, Kräne")
    add_footer(s, 4)
    notes(s, "~1 min: Kurz das System. Details in der Durchführung.")

    # 5 Soll-Ist
    s = blank(prs)
    add_bar(s)
    title_block(s, "Soll–Ist-Vergleich", "Freigabe versus heutiger Stand")
    table = s.shapes.add_table(6, 3, Inches(0.45), Inches(1.35), Inches(12.4), Inches(5.0)).table
    table_fill(
        table,
        ["Thema", "Soll (Freigabe)", "Ist (heute)"],
        [
            ["Linienfluss", "Roh → Bearbeitung → Montage → Palette → Lager", "Kunststoff E2E; Metall bis W2-SCL/FIO"],
            ["RFID", "Tag mit Artikel/Material/Typ", "4B Write + 4A Write; W2 ohne Reader 0"],
            ["Lager 54 Fächer", "Ein-/Auslagern, Suche, HMI", "W1 live; W2 SCL; OPC UA Gitter live"],
            ["HMI ein TP", "Manual, Auto, Meldungen, Regal", "Tags da; Simulation nur PLCSIM; Events offen"],
            ["Doku / Test", "Lasten-/Pflichtenheft, Test in FIO", "v1.6 / v1.3, Repo, Web OPC UA 192.168.0.1"],
        ],
        [Inches(2.4), Inches(5.0), Inches(5.0)],
    )
    add_footer(s, 5)
    notes(s, "~1 min: Ehrlich. Soll ist die Freigabe. Ist: Kunststoff Nachweis, Metall Spiegel, OPC UA live, HMI/CE offen.")

    # 6 Umfang
    s = blank(prs)
    add_bar(s)
    title_block(s, "Projektumfang", "Maschinengrenzen")
    card(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(5.0), "Gehört dazu", "Zonen 1A/1B bis 5A/5B\nSPS, HMI, Factory I/O, Sicherheitskonzept\nRFID, Vision, P&P, Palettierer, zwei Hochregale\nLagerverwaltung Online (Suche + OPC UA live)\nLastenheft, Pflichtenheft, SCL-Repo", head_color=GREEN)
    card(s, Inches(6.8), Inches(1.4), Inches(6.1), Inches(5.0), "Gehört nicht dazu", "Wasserverbrauch / Wasserwirtschaft\nERP-Anbindung\nexterne Energie, Stapler, Personal\nLive-Sync Web-DB ↔ SPS-Lager-DB\nvolle CE-Akte (später)", head_color=ACCENT)
    add_footer(s, 6)
    notes(s, "~50 s: Scope klar abgrenzen. Wasser und ERP bewusst draußen.")

    # 7 Kosten
    s = blank(prs)
    add_bar(s)
    title_block(s, "Kosten", "Schulungsprojekt: kein Capex für Mechanik — Aufwand = Zeit + vorhandene Lizenzen")
    table = s.shapes.add_table(6, 3, Inches(0.45), Inches(1.3), Inches(12.4), Inches(3.85)).table
    table_fill(
        table,
        ["Position", "Ansatz", "Bewertung"],
        [
            ["Mechanik / Schaltschrank", "0 EUR — Simulation Factory I/O", "keine Beschaffung"],
            ["SPS / HMI Hardware", "Labor (1518F, TP2200)", "Kurs / Schule, 0 EUR extra"],
            ["Software", "TIA V20, Factory I/O", "Schulungslizenz, im Kurs"],
            ["Arbeitszeit", "ca. 10 Wochen, ~250 h", "Steuerung, Test, Doku"],
            ["Zeitäquivalent*", "250 h × 45 EUR/h", "ca. 11.250 EUR intern"],
        ],
        [Inches(3.3), Inches(4.6), Inches(4.5)],
    )
    add_textbox(
        s,
        Inches(0.5),
        Inches(5.3),
        Inches(12.3),
        Inches(1.15),
        "*Interner Verrechnungssatz nur zur Einordnung (Annahme). Keine Rechnung an einen Kunden.\n"
        "Zusatzkosten über den Kurs hinaus: praktisch 0 EUR. Der Engpass war Kalenderzeit, nicht Geld.",
        size=14,
        color=GRAY,
    )
    add_footer(s, 7)
    notes(
        s,
        "~1 min: Klar sagen: keine Fabrik gebaut. 250 Stunden und 45 Euro sind Annahmen zum Einordnen. "
        "Kostenstelle ist die Weiterbildung, nicht ein Auftrag.",
    )

    # 8 Wirtschaftlichkeit
    s = blank(prs)
    add_bar(s)
    title_block(s, "Wirtschaftlichkeit", "Beispielrechnung für ein fiktives Unternehmen — Annahmen, keine echten Firmendaten")
    table = s.shapes.add_table(6, 3, Inches(0.45), Inches(1.25), Inches(12.4), Inches(3.55)).table
    table_fill(
        table,
        ["Annahme", "Ohne RFID / Inseln", "Mit dieser Lösung"],
        [
            ["Verwechslung Palette", "1 Fall / Tag", "0,1 Fall / Tag (Gate + Tag)"],
            ["Nacharbeit je Fall", "80 EUR", "80 EUR"],
            ["200 Arbeitstage / Jahr", "16.000 EUR Schaden", "1.600 EUR Rest"],
            ["Nutzen pro Jahr", "—", "ca. 14.400 EUR vermieden"],
            ["Interner Projektaufwand", "—", "ca. 11.250 EUR (250 h)"],
        ],
        [Inches(3.3), Inches(4.55), Inches(4.55)],
    )
    add_textbox(
        s,
        Inches(0.5),
        Inches(4.95),
        Inches(12.3),
        Inches(1.5),
        "Fazit der Rechnung: Amortisation unter einem Jahr — nur unter diesen Annahmen.\n"
        "Qualitativ (Freigabe): weniger Mischfehler Metall/Kunststoff, Übersicht Roh/Fertig, kürzere Durchlaufzeit,\n"
        "Nachverfolgbarkeit, digitales Abbild vor einem teuren Hallenumbau. Simulation spart Fehlversuche an der realen Anlage.",
        size=14,
        color=DARK,
    )
    add_footer(s, 8)
    notes(
        s,
        "~1,2 min: Zahlen laut ansagen als BEISPIEL. 1 Verwechslung/Tag mal 80 Euro mal 200 Tage. "
        "Nutzen 14400 gegen Aufwand 11250. Danach qualitativ: RFID-Gate, zwei getrennte Lager.",
    )

    # 9 Zeit
    s = blank(prs)
    add_bar(s)
    title_block(s, "Zeit", "Von der Freigabe bis zur Präsentation")
    table = s.shapes.add_table(6, 3, Inches(0.5), Inches(1.4), Inches(12.3), Inches(4.6)).table
    table_fill(
        table,
        ["Phase", "Zeitraum", "Ergebnis"],
        [
            ["Freigabe", "Juni 2026 (29.06.)", "Thema und Ziele bestätigt"],
            ["Konzept / Layout", "Juli 2026", "Zonen, Tags, W1/W2-Trennung"],
            ["Umsetzung Kern", "Juli–August 2026", "P&P, RFID, W1 Automatik E2E"],
            ["W2 + Doku + OPC UA", "August–September 2026", "W2 SCL, Lasten-/Pflichtenheft, Web live"],
            ["Stand heute", "19.09.2026", "FIO + OPC UA; HMI-Events und W2-Auto offen"],
        ],
        [Inches(3.0), Inches(3.5), Inches(5.8)],
    )
    add_footer(s, 9)
    notes(s, "~50 s: Zeitachse. Kunststoff zuerst war die bewusste Reihenfolge.")

    # 10 Herausforderungen
    s = blank(prs)
    add_bar(s)
    title_block(s, "Herausforderungen / Lösungen", "Was schiefgehen konnte — und wie es gelöst wurde")
    table = s.shapes.add_table(6, 2, Inches(0.5), Inches(1.35), Inches(12.3), Inches(5.1)).table
    table_fill(
        table,
        ["Herausforderung", "Lösung"],
        [
            ["W1 und W2 vermischt (gleiche FB-Namen)", "Eigene DBs, Merker %M60 / %M70, NW 29 / 28, Kräne 0 / 1"],
            ["HMI-DB hat RFID-Daten gelöscht", "Datenverwaltung kopiert Fachaktuell nur bei Fachwechsel"],
            ["PUT/GET liest keine optimized DBs", "Occupancy-Bytes HRL_Occ_* / HRL_2_Occ_* (%MB400 / %MB407)"],
            ["PC-WLAN ≠ CPU-Netz; HMI-Ethernet schreibt nicht", "PLCSIM Advanced 192.168.0.1; OPC UA 4840 live; HMI-Sim = PLCSIM"],
            ["F-CPU auf Standard-Sim-Instanz", "Neue Advanced-Instanz; CPU bleibt 1518F, kein 1518T"],
        ],
        [Inches(6.15), Inches(6.15)],
    )
    add_footer(s, 10)
    notes(s, "~1,2 min: Fünf echte Probleme. HMI-Ethernet auf dem Laptop ≠ echtes Panel.")

    # 11 Durchführung Überblick
    s = blank(prs)
    add_bar(s)
    title_block(s, "Projektdurchführung — Vorgehen", "Sechs Schritte aus der Freigabe")
    bullets(
        s,
        Inches(0.5),
        Inches(1.4),
        Inches(12.2),
        Inches(5.4),
        [
            "1. Layout, Signale, Betriebsarten Manual / Automatik",
            "2. SPS-Bausteine: Förder, Bearbeitung, RFID, P&P, Palettierer, Hochregal",
            "3. Lagerverwaltung: DBs, Suchen, Einlagern, Auslagern, Löschen",
            "4. HMI: Fach, Status, Meldungen, Regalbelegung",
            "5. Factory I/O: Bewegungen und Sensor/Aktor prüfen",
            "6. Tests, Fehlerreaktion, Dokumentation (Lasten- und Pflichtenheft)",
            "Arbeitsweise: zuerst Kunststofflinie komplett, dann Metall als Spiegel mit _W2",
        ],
        size=17,
    )
    add_footer(s, 11)
    notes(s, "~1 min: Methodik. Danach drei Folien zur konkreten Anlage.")

    # 12 Durchführung Linien
    s = blank(prs)
    add_bar(s)
    title_block(s, "Projektdurchführung — Anlage", "Zwei Linien, fünf Zonen")
    table = s.shapes.add_table(6, 3, Inches(0.45), Inches(1.35), Inches(12.4), Inches(5.1)).table
    table_fill(
        table,
        ["Zone", "Linie A Metall", "Linie B Kunststoff"],
        [
            ["1", "Roh, Roboter A/B, CNC", "Roh, Roboter C/D, CNC"],
            ["2", "Band + Vision", "Band + Vision → RFID-Daten"],
            ["3", "2-axis Montage", "2-axis Montage"],
            ["4", "3-axis P&P, Palettierer, RFID", "Palettierer, RFID Write"],
            ["5", "Hochregal W2  ·  NW 28", "Hochregal W1  ·  NW 29"],
        ],
        [Inches(1.5), Inches(5.45), Inches(5.45)],
    )
    add_footer(s, 12)
    notes(s, "~1,2 min: Spiegel. W1 = Kunststoff 5B. W2 = Metall 5A.")

    # 13 RFID Lager
    s = blank(prs)
    add_bar(s)
    title_block(s, "Projektdurchführung — RFID und Hochregal", "Identität bis ins Fach")
    bullets(
        s,
        Inches(0.5),
        Inches(1.4),
        Inches(12.2),
        Inches(5.4),
        [
            "Vision liefert Farbe, Artikel, Typ → RFID-CODE auf die Palette",
            "Kunststoff: Write wenn Palettierer fertig (Reader 2). Metall: Reader 1 in Zone 4A",
            "Gate: Palette vor dem Regal UND RFID erledigt — sonst kein Automatikzyklus",
            "W1 Automatik: State 10 Reader 5, dann freies Fach, Kran 0, Einlagern",
            "W2 Automatik: kein Reader 0 am Regal → State 10 entfällt; Produkt von 4A",
            "HMI: 6×9, Fach 1–6 unten; gelber Rand = Ziel-Fach",
            "Instanzen: fb_Datenverwaltung_Lager_DB  ≠  FB_Datenverwaltung_Lager_W2_DB",
        ],
        size=16,
    )
    add_footer(s, 13)
    notes(s, "~1,5 min: Wichtigste Technik-Folie. Gate und W1≠W2 betonen.")

    # 14 Software HMI
    s = blank(prs)
    add_bar(s)
    title_block(s, "Projektdurchführung — Software und HMI", "Ein Repo, ein OB1, ein Panel")
    card(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(4.9), "SPS / Repo", "Git: scl/ je Zone, docs/ je Kapitel\nOB1 29 Netzwerke + Plant Start/Stop\nSCL: Einlagern, Suchen, Automatik\nOPC UA: Si_Lagerverwaltung_Online live")
    card(s, Inches(6.8), Inches(1.4), Inches(6.1), Inches(4.9), "Bedienen", "TP2200 bleibt führend\nLaptop-Sim: HMI-Verbindung PLCSIM\nWeb: Suche SQLite + Live-Gitter OPC UA")
    add_footer(s, 14)
    notes(s, "~50 s: TP vor Web. OPC UA live. HMI-Sim über PLCSIM.")

    # 15 Ausblick
    s = blank(prs)
    add_bar(s)
    title_block(s, "Ausblick / Erweiterungen", "Was als Nächstes kommt")
    bullets(
        s,
        Inches(0.5),
        Inches(1.4),
        Inches(12.2),
        Inches(5.4),
        [
            "Warehouse 2: erste Palette Hand und Automatik (ohne Reader 0)",
            "Zone 1/2: CNC- und Band-Feinsteuerung in OB1 festziehen",
            "WinCC: Start/Stop %M58 verdrahten, 54-Fach-Animation W1/W2",
            "Plant-FB in TIA als FUP zeichnen (Rezept liegt vor)",
            "Sicherheit nach EN ISO 12100 und CE-Dokumentation",
            "Nicht geplant: ERP, Wasserwirtschaft, reale Mechanik, 1518T",
        ],
        size=17,
    )
    add_footer(s, 15)
    notes(s, "~1 min: Ausblick konkret. OPC UA ist erledigt — HMI und W2-Auto bleiben.")

    # 16 Fazit
    s = blank(prs)
    add_bar(s)
    title_block(s, "Fazit — was ich daraus gelernt habe", "")
    bullets(
        s,
        Inches(0.5),
        Inches(1.35),
        Inches(12.2),
        Inches(5.5),
        [
            "Eine Factory braucht klare Grenzen: Wer ist W1, wer ist W2 — sonst zerstört Copy-Paste die Daten",
            "RFID nützt nur, wenn das Gate den Prozess wirklich sperrt",
            "HMI-DBs und Prozess-DBs trennen: jedes Scan kopieren kann Buchungen löschen",
            "Kommunikation SPS–PC: PLCSIM Advanced und OPC UA — nicht Classic, nicht Laptop-WLAN",
            "Lastenheft (Was) und Pflichtenheft (Wie) früh schreiben — sonst driftet der Code",
            "Zuerst eine Linie fertig machen (Kunststoff), dann spiegeln — schneller als beide halb",
        ],
        size=17,
    )
    add_footer(s, 16)
    notes(s, "~1,2 min: Lernen, nicht noch einmal die Zonen aufzählen.")

    # 17 Kernaussage
    s = blank(prs)
    add_bar(s)
    title_block(s, "Kernaussage", "In einem Satz")
    add_textbox(
        s,
        Inches(0.7),
        Inches(2.2),
        Inches(12),
        Inches(2.8),
        "Zwei Linien, RFID bis ins Fach,\nWarehouse 1 Kunststoff und Warehouse 2 Metall\ngetrennt, ein HMI — nachgewiesen in TIA und Factory I/O.",
        size=26,
        bold=True,
        color=NAVY,
        align=PP_ALIGN.CENTER,
    )
    add_footer(s, 17)
    notes(s, "~20 s: Satz stehen lassen. Dann Folie Ende.")

    # 18 Ende
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), SLIDE_H)
    acc.fill.solid()
    acc.fill.fore_color.rgb = TEAL
    acc.line.fill.background()
    add_textbox(s, Inches(0.7), Inches(2.3), Inches(12), Inches(0.9), "Vielen Dank für Ihre Aufmerksamkeit", size=32, bold=True, color=WHITE)
    add_textbox(s, Inches(0.7), Inches(3.5), Inches(12), Inches(0.6), "Fragen?", size=28, color=TEAL)
    add_textbox(
        s,
        Inches(0.7),
        Inches(4.6),
        Inches(12),
        Inches(1.2),
        "Dereje Hailemariam  ·  Berlin  ·  19.09.2026\nWeiterbildung Abschlussprojekt Automatisierungstechnik",
        size=16,
        color=WHITE,
    )
    notes(s, "~45 s: Danke, Fragen. Optional Demo FIO oder Streamlit.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    target = OUT
    try:
        prs.save(target)
    except PermissionError:
        target = OUT.with_name(OUT.stem + "_neu.pptx")
        prs.save(target)
        print("locked:", OUT.name, "-> wrote", target.name)
    print("wrote", target, "slides", len(prs.slides))


if __name__ == "__main__":
    build()
